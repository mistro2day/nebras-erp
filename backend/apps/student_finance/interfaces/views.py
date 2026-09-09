from rest_framework import status, filters
from rest_framework.decorators import action
from rest_framework.response import Response
from django.utils import timezone
from django.db.models import Sum, Count, Q

from apps.shared.interfaces.views import BaseCRUDViewSet
from apps.common.responses import StandardResponse

from apps.student_finance.domain.models import (
    FeeCategory, FeeType, FeeStructure, FeeSchedule, AcademicFeePlan,
    StudentBillingAccount, StudentInvoice, InvoiceItem, InvoiceAdjustment,
    InvoiceDiscount, Scholarship, ScholarshipRule, FinancialAid,
    InstallmentPlan, Installment, StudentReceivable, PaymentAllocation,
    Receipt, Refund, CreditNote, DebitNote, LateFeeRule, CollectionPolicy,
    FinancialHold, BillingCycle, Statement, BillingAudit, StudentFinanceSettings
)

from apps.student_finance.interfaces.serializers import (
    FeeCategorySerializer, FeeTypeSerializer, FeeStructureSerializer,
    FeeScheduleSerializer, AcademicFeePlanSerializer, StudentBillingAccountSerializer,
    StudentInvoiceSerializer, InvoiceItemSerializer, InvoiceAdjustmentSerializer,
    InvoiceDiscountSerializer, ScholarshipSerializer, ScholarshipRuleSerializer,
    FinancialAidSerializer, InstallmentPlanSerializer, InstallmentSerializer,
    StudentReceivableSerializer, PaymentAllocationSerializer, ReceiptSerializer,
    RefundSerializer, CreditNoteSerializer, DebitNoteSerializer,
    LateFeeRuleSerializer, CollectionPolicySerializer, FinancialHoldSerializer,
    BillingCycleSerializer, StatementSerializer, BillingAuditSerializer,
    StudentFinanceSettingsSerializer
)

from apps.student_finance.application.services import (
    BillingService, PaymentService, ScholarshipService, HoldService
)
from apps.student_finance.application import online_payment as online_payment_service
from apps.student_finance.domain.models import OnlinePaymentRequest
from apps.student_finance.interfaces.serializers import OnlinePaymentRequestSerializer
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from django.core.exceptions import ValidationError as DjangoValidationError


class FeeCategoryViewSet(BaseCRUDViewSet):
    model_class = FeeCategory
    serializer_class = FeeCategorySerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name_ar', 'name_en', 'code']


class FeeTypeViewSet(BaseCRUDViewSet):
    model_class = FeeType
    serializer_class = FeeTypeSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name_ar', 'name_en', 'code']


class FeeStructureViewSet(BaseCRUDViewSet):
    model_class = FeeStructure
    serializer_class = FeeStructureSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['name', 'academic_year']


class FeeScheduleViewSet(BaseCRUDViewSet):
    model_class = FeeSchedule
    serializer_class = FeeScheduleSerializer


class AcademicFeePlanViewSet(BaseCRUDViewSet):
    model_class = AcademicFeePlan
    serializer_class = AcademicFeePlanSerializer


class StudentBillingAccountViewSet(BaseCRUDViewSet):
    model_class = StudentBillingAccount
    serializer_class = StudentBillingAccountSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['account_number']

    @action(detail=False, methods=['get'], url_path='dashboard-stats')
    def get_dashboard_stats(self, request):
        """جلب إحصائيات لوحة التحكم المالية للطلاب."""
        tenant_id = request.tenant_id
        
        # إجمالي حسابات القبض المستحقة
        outstanding_receivables = StudentReceivable.objects.filter(
            tenant_id=tenant_id, status='outstanding'
        ).aggregate(total=Sum('outstanding_amount'))['total'] or 0.0

        # تحصيلات اليوم
        today = timezone.localdate()
        today_collections = Receipt.objects.filter(
            tenant_id=tenant_id, payment_date=today, status='posted'
        ).aggregate(total=Sum('amount'))['total'] or 0.0

        # تحصيلات الشهر
        start_of_month = today.replace(day=1)
        monthly_collections = Receipt.objects.filter(
            tenant_id=tenant_id, payment_date__gte=start_of_month, status='posted'
        ).aggregate(total=Sum('amount'))['total'] or 0.0

        # الحظر المالي النشط
        active_holds = FinancialHold.objects.filter(
            tenant_id=tenant_id, status='active'
        ).count()

        # طلبات الاسترداد المعلقة
        pending_refunds = Refund.objects.filter(
            tenant_id=tenant_id, status='pending'
        ).count()

        # الأقساط المستحقة خلال 7 أيام القادمة
        next_week = today + timezone.timedelta(days=7)
        due_installments = Installment.objects.filter(
            tenant_id=tenant_id, due_date__range=(today, next_week), status='pending'
        ).aggregate(total=Sum('amount'))['total'] or 0.0

        stats = {
            'outstanding_receivables': float(outstanding_receivables),
            'today_collections': float(today_collections),
            'monthly_collections': float(monthly_collections),
            'active_holds': active_holds,
            'pending_refunds': pending_refunds,
            'due_installments': float(due_installments)
        }
        return Response(stats, status=status.HTTP_200_OK)

    @action(detail=True, methods=['get'], url_path='statement')
    def statement(self, request, pk=None):
        """كشف حساب مالي تفصيلي شامل للطالب مع تفاصيل وهوية المستأجر والبحث المرن."""
        tenant = getattr(request, 'tenant', None)
        tenant_id = request.tenant_id

        # 0. البحث المرن عن الحساب المالي (بمعرف الحساب، أو معرف الطالب، أو رقم الحساب)
        account = None
        if pk:
            # محاولة كـ UUID لحساب الفوترة
            account = StudentBillingAccount.objects.filter(tenant_id=tenant_id, id=pk).first()
            if not account:
                # محاولة كـ student_id
                account = StudentBillingAccount.objects.filter(tenant_id=tenant_id, student_id=pk).first()
            if not account:
                # محاولة كـ account_number
                account = StudentBillingAccount.objects.filter(tenant_id=tenant_id, account_number=pk).first()

        # إذا لم يوجد حساب فوترة، نبحث عن الطالب وننشئ له حساباً تلقائياً
        if not account and pk:
            from apps.students.domain.models import Student
            student_obj = Student.objects.filter(tenant_id=tenant_id, id=pk).first()
            if not student_obj:
                student_obj = Student.objects.filter(tenant_id=tenant_id, student_number=pk).first()
            
            if student_obj:
                from decimal import Decimal
                account = StudentBillingAccount.objects.create(
                    tenant_id=tenant_id,
                    student_id=student_obj.id,
                    account_number=f"ACC-ST-{timezone.now().strftime('%y%m%d%H%M')}-{student_obj.student_number}",
                    opening_balance=Decimal('0.0'),
                    current_balance=Decimal('0.0'),
                    outstanding_balance=Decimal('0.0'),
                    credit_balance=Decimal('0.0'),
                    created_by=request.user.id if (request.user and request.user.is_authenticated) else None
                )

        if not account:
            return StandardResponse(
                message="لم يتم العثور على حساب مالي أو طالب مطابق لهذا المعرف.",
                success=False,
                status=status.HTTP_404_NOT_FOUND
            )

        from apps.student_finance.interfaces.serializers import _extract_student_finance_metadata
        student_meta = _extract_student_finance_metadata(account)

        # 1. بيانات وهوية المستأجر (المدرسة)
        tenant = getattr(request, 'tenant', None)
        if not tenant and hasattr(account, 'tenant_id') and account.tenant_id:
            from apps.tenants.models import Tenant
            tenant = Tenant.objects.filter(id=account.tenant_id).first()
        if not tenant:
            from apps.tenants.models import Tenant
            tenant = Tenant.objects.first()

        tenant_name_ar = (tenant.name_ar or tenant.name) if tenant else 'مدارس المورد النموذجية الخاصة'
        tenant_name_en = (tenant.name_en or tenant.name) if (tenant and tenant.name_en) else 'Al-Mawred Model Private Schools'

        tenant_info = {
            'id': str(tenant.id) if tenant else str(tenant_id),
            'name': getattr(tenant, 'name', '') or tenant_name_ar,
            'name_ar': tenant_name_ar,
            'name_en': tenant_name_en,
            'logo_url': request.build_absolute_uri(tenant.logo.url) if (tenant and getattr(tenant, 'logo', None)) else '/assets/branding/al_mawadda_logo.jpg',
            'stamp_url': request.build_absolute_uri(tenant.stamp.url) if (tenant and getattr(tenant, 'stamp', None)) else '',
            'phone': getattr(tenant, 'phone_number', '') or '09123456789',
            'email': getattr(tenant, 'email', '') or 'accounts@almawred.edu.sd',
            'address': getattr(tenant, 'address', '') or 'الخرطوم - العمارات - شارع 15، السودان',
            'affiliation': 'المرحلة الابتدائية والمتوسطة والثانوية • ولاية الخرطوم',
        }

        # 2. بيانات الحساب
        account_info = {
            'id': str(account.id),
            'account_number': account.account_number,
            'opening_balance': float(account.opening_balance or 0.0),
            'current_balance': float(account.current_balance or 0.0),
            'outstanding_balance': float(account.outstanding_balance or 0.0),
            'credit_balance': float(account.credit_balance or 0.0),
            'financial_hold': account.financial_hold,
            'currency': 'ج.س',
        }

        # 3. تجميع الحركات المالية وترتيبها زمنياً لحساب الرصيد التراكمي
        transactions = []

        # رصيد افتتاحي
        if account.opening_balance and float(account.opening_balance) > 0:
            transactions.append({
                'date': str(account.created_at.date() if account.created_at else timezone.localdate()),
                'type': 'opening_balance',
                'type_label': 'رصيد افتتاحي سابق',
                'reference_number': account.account_number,
                'description': 'الرصيد الافتتاحي المقيد عند فتح الحساب المالي',
                'debit': float(account.opening_balance),
                'credit': 0.0,
                'payment_method': '',
            })

        # فواتير الرسوم
        invoices = StudentInvoice.objects.filter(
            student_billing_account=account, tenant_id=tenant_id
        ).order_by('issue_date', 'created_at')
        
        for inv in invoices:
            transactions.append({
                'date': str(inv.issue_date or inv.created_at.date()),
                'type': 'invoice',
                'type_label': 'فاتورة رسوم دراسية',
                'reference_number': inv.invoice_number,
                'description': f"فاتورة رسوم - تاريخ الاستحقاق: {inv.due_date}",
                'debit': float(inv.total_amount or 0.0),
                'credit': 0.0,
                'payment_method': '',
                'invoice_status': inv.status,
            })

        # سندات القبض والتحصيلات
        receipts = Receipt.objects.filter(
            student_billing_account=account, tenant_id=tenant_id
        ).order_by('payment_date', 'created_at')

        # جلب مسميات طرق الدفع المعتمدة
        payment_methods_map = {}
        try:
            from apps.finance.domain.models import PaymentMethod
            pms = PaymentMethod.objects.filter(tenant_id=tenant_id)
            payment_methods_map = {pm.id: (pm.name_ar or pm.name) for pm in pms}
        except Exception:
            pass

        for rec in receipts:
            pm_name = payment_methods_map.get(rec.payment_method_id, 'سداد نقدي / بنكك')
            transactions.append({
                'date': str(rec.payment_date or (rec.created_at.date() if rec.created_at else timezone.localdate())),
                'type': 'receipt',
                'type_label': 'سند قبض / تحصيل',
                'reference_number': rec.receipt_number,
                'description': f"سداد رسوم دراسية - إيصال رقم {rec.receipt_number}",
                'debit': 0.0,
                'credit': float(rec.amount or 0.0),
                'payment_method': pm_name,
                'reference_trans': str(rec.voucher_id or rec.receipt_number or ''),
            })

        # المنح والخصومات
        scholarships = Scholarship.objects.filter(
            student_billing_account=account, tenant_id=tenant_id
        ).order_by('start_date', 'created_at')

        for sch in scholarships:
            val = float(sch.fixed_amount or 0.0)
            if val > 0:
                transactions.append({
                    'date': str(sch.start_date or sch.created_at.date()),
                    'type': 'scholarship',
                    'type_label': f"منحة ({sch.name})",
                    'reference_number': sch.code or 'SCH',
                    'description': f"منحة دراسية معتمدة: {sch.name}",
                    'debit': 0.0,
                    'credit': val,
                    'payment_method': '',
                })

        # فرز الحركات تصاعدياً بالتاريخ
        transactions.sort(key=lambda x: x['date'])

        # حساب الرصيد التراكمي (Running Balance)
        running_bal = 0.0
        total_invoiced = 0.0
        total_paid = 0.0
        total_discounted = 0.0

        for t in transactions:
            debit = t['debit']
            credit = t['credit']
            running_bal += (debit - credit)
            t['running_balance'] = float(running_bal)
            
            if t['type'] in ('invoice', 'opening_balance'):
                total_invoiced += debit
            if t['type'] == 'receipt':
                total_paid += credit
            if t['type'] == 'scholarship':
                total_discounted += credit

        # جدول الأقساط الحقيقية للطالب
        installments_qs = list(Installment.objects.filter(
            student_billing_account=account, tenant_id=tenant_id
        ).order_by('due_date'))

        total_inst = len(installments_qs)
        ordinal_ar = {1: 'الأول', 2: 'الثاني', 3: 'الثالث', 4: 'الرابع', 5: 'الخامس', 6: 'السادس'}

        installments_data = []
        for idx, ins in enumerate(installments_qs, start=1):
            amt = float(ins.amount or 0.0)
            paid = float(ins.paid_amount or 0.0)
            rem = max(0.0, amt - paid)

            seq = ordinal_ar.get(idx, f"رقم {idx}")
            if total_inst == 1:
                label = "الدفعة الكاملة (100%)"
            elif total_inst == 2:
                label = f"القسط {seq} (50%)"
            elif total_inst == 3:
                pct = "40%" if idx == 1 else "30%"
                label = f"القسط {seq} ({pct})"
            elif total_inst == 4:
                label = f"القسط {seq} (25%)"
            else:
                label = f"القسط {seq} من {total_inst}"

            plan_name = ins.installment_plan.name if ins.installment_plan else 'خطة السداد بالأقساط'
            full_title = f"{label} — {plan_name}"

            installments_data.append({
                'id': str(ins.id),
                'due_date': str(ins.due_date),
                'amount': amt,
                'paid_amount': paid,
                'remaining_amount': rem,
                'status': ins.status,
                'status_label': 'مسدد بالكامل' if ins.status == 'paid' else ('متأخر' if str(ins.due_date) < str(timezone.localdate()) else 'مجدول'),
                'plan_name': full_title,
                'installment_label': label,
            })

        # في حال عدم وجود أقساط مسجلة ولكن توجد فواتير برصيد متبقٍ أو مسددة
        if total_inst == 0 and invoices.exists():
            for inv in invoices:
                tot = float(inv.total_amount or 0.0)
                pd = float(inv.paid_amount or 0.0)
                rem = float(inv.outstanding_amount if inv.outstanding_amount is not None else max(0.0, tot - pd))
                due_d = str(inv.due_date or inv.issue_date or timezone.localdate())
                is_paid = (rem <= 0)
                is_overdue = (not is_paid and due_d < str(timezone.localdate()))
                installments_data.append({
                    'id': str(inv.id),
                    'due_date': due_d,
                    'amount': tot,
                    'paid_amount': pd,
                    'remaining_amount': rem,
                    'status': 'paid' if is_paid else ('overdue' if is_overdue else 'pending'),
                    'status_label': 'مسدد بالكامل' if is_paid else ('متأخر' if is_overdue else 'مجدول'),
                    'plan_name': f"استحقاق فاتورة {inv.invoice_number}",
                    'installment_label': f"فاتورة {inv.invoice_number}",
                })

        import uuid
        statement_number = f"STMT-{timezone.localdate().strftime('%Y%m%d')}-{account.account_number[-4:] if len(account.account_number) >= 4 else uuid.uuid4().hex[:4].upper()}"

        return StandardResponse(data={
            'meta': {
                'statement_number': statement_number,
                'generated_at': timezone.now().strftime('%Y-%m-%d %H:%M'),
                'academic_year': '2026 / 2027',
                'currency': 'ج.س',
            },
            'tenant': tenant_info,
            'student': student_meta,
            'account': account_info,
            'summary': {
                'total_invoiced': float(total_invoiced),
                'total_paid': float(total_paid),
                'total_discounted': float(total_discounted),
                'net_outstanding': float(account.outstanding_balance or running_bal),
                'credit_balance': float(account.credit_balance or 0.0),
                'installments_count': len(installments_data),
                'installments_total': sum(i['amount'] for i in installments_data),
                'installments_paid': sum(i['paid_amount'] for i in installments_data),
                'installments_remaining': sum(i['remaining_amount'] for i in installments_data),
            },
            'transactions': transactions,
            'installments': installments_data,
        })


class StudentInvoiceViewSet(BaseCRUDViewSet):
    model_class = StudentInvoice
    serializer_class = StudentInvoiceSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['invoice_number']

    def get_queryset(self):
        qs = super().get_queryset()
        account = self.request.query_params.get('student_billing_account')
        status_param = self.request.query_params.get('status')
        if account:
            qs = qs.filter(student_billing_account_id=account)
        if status_param:
            qs = qs.filter(status=status_param)
        return qs

    @action(detail=False, methods=['post'], url_path='generate-invoice')
    def generate_invoice(self, request):
        tenant_id = request.tenant_id
        billing_account_id = request.data.get('billing_account_id')
        fee_structure_ids = request.data.get('fee_structure_ids', [])
        due_date_str = request.data.get('due_date')
        
        if not billing_account_id or not fee_structure_ids or not due_date_str:
            return Response({'error': 'billing_account_id, fee_structure_ids, and due_date are required'}, status=status.HTTP_400_BAD_REQUEST)

        fee_structures = FeeStructure.objects.filter(id__in=fee_structure_ids, tenant_id=tenant_id)
        due_date = timezone.datetime.strptime(due_date_str, '%Y-%m-%d').date()

        invoice = BillingService.generate_student_invoice(
            tenant_id=tenant_id,
            billing_account_id=billing_account_id,
            fee_structures=fee_structures,
            due_date=due_date,
            user_id=request.user.id if request.user else None
        )
        serializer = self.get_serializer(invoice)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class InvoiceItemViewSet(BaseCRUDViewSet):
    model_class = InvoiceItem
    serializer_class = InvoiceItemSerializer


class InvoiceAdjustmentViewSet(BaseCRUDViewSet):
    model_class = InvoiceAdjustment
    serializer_class = InvoiceAdjustmentSerializer


class InvoiceDiscountViewSet(BaseCRUDViewSet):
    model_class = InvoiceDiscount
    serializer_class = InvoiceDiscountSerializer


class ScholarshipViewSet(BaseCRUDViewSet):
    model_class = Scholarship
    serializer_class = ScholarshipSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        account = self.request.query_params.get('student_billing_account')
        if account:
            qs = qs.filter(student_billing_account_id=account)
        return qs

    @action(detail=False, methods=['post'], url_path='apply-scholarship')
    def apply_scholarship(self, request):
        tenant_id = request.tenant_id
        billing_account_id = request.data.get('billing_account_id')
        name = request.data.get('name')
        scholarship_type = request.data.get('type')
        amount_percentage = request.data.get('amount_percentage', 0.0)
        fixed_amount = request.data.get('fixed_amount', 0.0)
        start_date = request.data.get('start_date')
        end_date = request.data.get('end_date')

        sc = ScholarshipService.apply_scholarship(
            tenant_id=tenant_id,
            billing_account_id=billing_account_id,
            name=name,
            scholarship_type=scholarship_type,
            amount_percentage=amount_percentage,
            fixed_amount=fixed_amount,
            start_date=start_date,
            end_date=end_date,
            user_id=request.user.id if request.user else None
        )
        serializer = self.get_serializer(sc)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class ScholarshipRuleViewSet(BaseCRUDViewSet):
    model_class = ScholarshipRule
    serializer_class = ScholarshipRuleSerializer


class FinancialAidViewSet(BaseCRUDViewSet):
    model_class = FinancialAid
    serializer_class = FinancialAidSerializer


class InstallmentPlanViewSet(BaseCRUDViewSet):
    model_class = InstallmentPlan
    serializer_class = InstallmentPlanSerializer


def _sync_invoices_to_installments(tenant_id):
    """
    مزامنة آلية لأي فواتير صادرة للطلاب بدون أقساط مجدولة، لضمان إدراجها
    فورياً في تقويم الدفعات والأقساط وتصدير المستحقين الشهري.
    """
    if not tenant_id:
        return
    # فحص سريع وخفيف أولاً دون تحميل كل السجلات لتفادي إبطاء الاستجابة
    if not StudentInvoice.objects.filter(tenant_id=tenant_id, installments__isnull=True).exists():
        return
    invoices = list(StudentInvoice.objects.filter(tenant_id=tenant_id, installments__isnull=True))
    if not invoices:
        return

    plan = InstallmentPlan.objects.filter(tenant_id=tenant_id, is_active=True).first()
    if not plan:
        plan = InstallmentPlan.objects.create(
            tenant_id=tenant_id,
            name='خطة الأقساط المعتمدة',
            number_of_installments=1,
            grace_period_days=7,
            is_active=True
        )

    for inv in invoices:
        due = inv.due_date or inv.issue_date or timezone.localdate()
        amt = inv.total_amount or 0
        paid = inv.paid_amount or 0
        is_paid = (paid >= amt and amt > 0)
        Installment.objects.create(
            tenant_id=tenant_id,
            student_billing_account=inv.student_billing_account,
            invoice=inv,
            installment_plan=plan,
            due_date=due,
            amount=amt,
            paid_amount=paid,
            status='paid' if is_paid else 'pending',
            created_by=inv.created_by
        )


class InstallmentViewSet(BaseCRUDViewSet):
    model_class = Installment
    serializer_class = InstallmentSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['invoice__invoice_number', 'student_billing_account__account_number']

    def get_queryset(self):
        tenant_id = getattr(self.request, 'tenant_id', None)
        if tenant_id:
            _sync_invoices_to_installments(tenant_id)

        qs = super().get_queryset()
        account = self.request.query_params.get('student_billing_account')
        status_param = self.request.query_params.get('status')
        month = self.request.query_params.get('month')
        year = self.request.query_params.get('year')
        range_param = self.request.query_params.get('range')
        
        today = timezone.localdate()
        if account:
            qs = qs.filter(student_billing_account_id=account)
        if status_param and status_param != 'all':
            if status_param == 'overdue':
                qs = qs.filter(due_date__lt=today).exclude(status='paid')
            else:
                qs = qs.filter(status=status_param)
        if year and month:
            try:
                qs = qs.filter(due_date__year=int(year), due_date__month=int(month))
            except ValueError:
                pass
        if range_param == 'today':
            qs = qs.filter(due_date=today)
        elif range_param == 'this_week':
            week_end = today + timezone.timedelta(days=7)
            qs = qs.filter(due_date__range=(today, week_end))
        elif range_param == 'overdue':
            qs = qs.filter(due_date__lt=today).exclude(status='paid')

        return qs.order_by('due_date', 'amount')

    @action(detail=False, methods=['get'], url_path='calendar')
    def calendar(self, request):
        """عرض تقويم استحقاق الأقساط والدفعات للطلاب مع مؤشرات ذكية وتصنيف زمني."""
        tenant_id = getattr(request, 'tenant_id', None)
        if not tenant_id and hasattr(request, 'tenant') and request.tenant:
            tenant_id = request.tenant.id
        if not tenant_id:
            from apps.tenants.domain.models import Tenant
            t_obj = Tenant.objects.filter(is_active=True).first()
            if t_obj:
                tenant_id = t_obj.id

        today = timezone.localdate()
        if tenant_id:
            _sync_invoices_to_installments(tenant_id)
        
        # قراءة معايير التاريخ
        try:
            year = int(request.query_params.get('year', today.year))
            month = int(request.query_params.get('month', today.month))
        except (ValueError, TypeError):
            year, month = today.year, today.month

        status_param = request.query_params.get('status', 'all')
        range_param = request.query_params.get('range', 'all')
        search_query = request.query_params.get('search', '').strip()

        # استعلام الأساس
        base_qs = Installment.objects.filter(tenant_id=tenant_id)
        
        # مؤشرات عامة سريعة (KPI Counters) مدمجة في استعلام واحد فائق السرعة
        next_7_days = today + timezone.timedelta(days=7)
        kpi = base_qs.aggregate(
            overdue_count=Count('id', filter=Q(due_date__lt=today) & ~Q(status='paid')),
            overdue_amount=Sum('amount', filter=Q(due_date__lt=today) & ~Q(status='paid')),
            due_today_count=Count('id', filter=Q(due_date=today) & ~Q(status='paid')),
            due_today_amount=Sum('amount', filter=Q(due_date=today) & ~Q(status='paid')),
            due_week_count=Count('id', filter=Q(due_date__range=(today, next_7_days)) & ~Q(status='paid')),
            due_week_amount=Sum('amount', filter=Q(due_date__range=(today, next_7_days)) & ~Q(status='paid')),
            paid_month_count=Count('id', filter=Q(due_date__year=year, due_date__month=month, status='paid')),
            paid_month_amount=Sum('amount', filter=Q(due_date__year=year, due_date__month=month, status='paid')),
        )
        overdue_count = kpi.get('overdue_count') or 0
        overdue_amount = float(kpi.get('overdue_amount') or 0.0)
        due_today_count = kpi.get('due_today_count') or 0
        due_today_amount = float(kpi.get('due_today_amount') or 0.0)
        due_week_count = kpi.get('due_week_count') or 0
        due_week_amount = float(kpi.get('due_week_amount') or 0.0)
        paid_month_count = kpi.get('paid_month_count') or 0
        paid_month_amount = float(kpi.get('paid_month_amount') or 0.0)

        # التصفية حسب نطاق الشهر المطلوب للتقويم
        month_qs = base_qs.filter(due_date__year=year, due_date__month=month)
        
        # التجميع اليومي للشهر المختار (Calendar Matrix)
        days_summary = {}
        for item in month_qs.values('due_date', 'status').annotate(total_amount=Sum('amount'), count=Count('id')):
            d_str = str(item['due_date'])
            if d_str not in days_summary:
                days_summary[d_str] = {
                    'date': d_str,
                    'count': 0,
                    'total_amount': 0.0,
                    'pending_count': 0,
                    'paid_count': 0,
                    'overdue_count': 0,
                    'has_overdue': False,
                    'is_today': (item['due_date'] == today)
                }
            st = item['status']
            amt = float(item['total_amount'] or 0.0)
            c = int(item['count'] or 1)
            days_summary[d_str]['count'] += c
            days_summary[d_str]['total_amount'] += amt
            if st == 'paid':
                days_summary[d_str]['paid_count'] += c
            elif item['due_date'] < today and st != 'paid':
                days_summary[d_str]['overdue_count'] += c
                days_summary[d_str]['has_overdue'] = True
            else:
                days_summary[d_str]['pending_count'] += c

        # قائمة الأقساط المعروضة حسب الفلتر
        filtered_qs = base_qs
        if range_param == 'today':
            filtered_qs = filtered_qs.filter(due_date=today)
        elif range_param == 'this_week':
            filtered_qs = filtered_qs.filter(due_date__range=(today, next_7_days))
        elif range_param == 'overdue':
            filtered_qs = filtered_qs.filter(due_date__lt=today).exclude(status='paid')
        elif range_param == 'this_month':
            filtered_qs = filtered_qs.filter(due_date__year=year, due_date__month=month)
        else:
            filtered_qs = filtered_qs.filter(due_date__year=year, due_date__month=month)

        if status_param and status_param != 'all':
            if status_param == 'overdue':
                filtered_qs = filtered_qs.filter(due_date__lt=today).exclude(status='paid')
            else:
                filtered_qs = filtered_qs.filter(status=status_param)

        filtered_qs = filtered_qs.select_related('student_billing_account', 'invoice', 'installment_plan').order_by('due_date', 'status')
        installments_list = list(filtered_qs[:200])

        # تحميل بيانات الطلاب المسبقة دفعة واحدة لمنع استعلامات N+1 وتسريع الاستجابة
        student_ids = {
            ins.student_billing_account.student_id
            for ins in installments_list
            if ins.student_billing_account and ins.student_billing_account.student_id
        }
        if student_ids:
            from apps.students.domain.models import Student
            from apps.academics.domain.models import Grade, Section
            from apps.student_finance.interfaces.serializers import _extract_student_finance_metadata
            students_qs = Student.objects.filter(id__in=student_ids).select_related('profile').prefetch_related(
                'enrollments', 'family_relations'
            )
            student_map = {s.id: s for s in students_qs}

            grade_ids = set()
            section_ids = set()
            for s in students_qs:
                for e in s.enrollments.all():
                    if getattr(e, 'grade_id', None):
                        grade_ids.add(e.grade_id)
                    if getattr(e, 'section_id', None):
                        section_ids.add(e.section_id)

            grade_map = {g.id: (getattr(g, 'name_ar', '') or getattr(g, 'name', '')) for g in Grade.objects.filter(id__in=grade_ids)} if grade_ids else {}
            section_map = {sec.id: (getattr(sec, 'name_ar', '') or getattr(sec, 'name', '')) for sec in Section.objects.filter(id__in=section_ids)} if section_ids else {}

            for ins in installments_list:
                if ins.student_billing_account:
                    _extract_student_finance_metadata(
                        ins.student_billing_account,
                        student_map=student_map,
                        grade_map=grade_map,
                        section_map=section_map
                    )

        serialized_data = InstallmentSerializer(installments_list, many=True).data

        if search_query:
            q_lower = search_query.lower()
            serialized_data = [
                d for d in serialized_data
                if q_lower in (d.get('student_name') or '').lower()
                or q_lower in (d.get('student_number') or '').lower()
                or q_lower in (d.get('invoice_number') or '').lower()
                or q_lower in (d.get('guardian_phone') or '').lower()
            ]

        for d in serialized_data:
            due_d = d.get('due_date')
            if due_d and d.get('status') != 'paid' and str(due_d) < str(today):
                d['computed_status'] = 'overdue'
                d['computed_status_label'] = 'متأخر السداد'
            elif due_d and str(due_d) == str(today) and d.get('status') != 'paid':
                d['computed_status'] = 'due_today'
                d['computed_status_label'] = 'يستحق اليوم'
            elif d.get('status') == 'paid':
                d['computed_status'] = 'paid'
                d['computed_status_label'] = 'مسدد بالكامل'
            else:
                d['computed_status'] = 'pending'
                d['computed_status_label'] = 'مجدول'

        return StandardResponse(data={
            'meta': {
                'year': year,
                'month': month,
                'today': str(today),
            },
            'summary': {
                'overdue_count': overdue_count,
                'overdue_amount': overdue_amount,
                'due_today_count': due_today_count,
                'due_today_amount': due_today_amount,
                'due_week_count': due_week_count,
                'due_week_amount': due_week_amount,
                'paid_month_count': paid_month_count,
                'paid_month_amount': paid_month_amount,
                'month_total_due': float(month_qs.exclude(status='paid').aggregate(s=Sum('amount'))['s'] or 0.0),
            },
            'days_summary': days_summary,
            'installments': serialized_data
        })

    @action(detail=True, methods=['post'], url_path='quick-pay')
    def quick_pay(self, request, pk=None):
        """تسجيل سداد فوري لقسط مجدول وتحديث الفاتورة وإصدار سند قبض."""
        tenant_id = request.tenant_id
        installment = self.get_object()
        
        remaining = max(0.0, float(installment.amount) - float(installment.paid_amount))
        amount_to_pay = float(request.data.get('amount') or remaining)
        if amount_to_pay <= 0:
            return StandardResponse(message="مبلغ السداد يجب أن يكون أكبر من الصفر.", success=False, status=status.HTTP_400_BAD_REQUEST)
        
        reference_number = request.data.get('reference_number', '')
        payment_method_name = request.data.get('payment_method', 'تطبيق بنكك (بنك الخرطوم)')
        notes = request.data.get('notes', 'سداد قسط دراسي عبر تقويم الدفعات')
        
        new_paid = float(installment.paid_amount) + amount_to_pay
        installment.paid_amount = min(new_paid, float(installment.amount))
        if installment.paid_amount >= float(installment.amount):
            installment.status = 'paid'
        installment.save()

        if installment.invoice:
            inv = installment.invoice
            inv.paid_amount = float(inv.paid_amount or 0.0) + amount_to_pay
            inv.outstanding_amount = max(0.0, float(inv.total_amount) - float(inv.paid_amount))
            inv.save()

        import uuid
        receipt_no = f"REC-{timezone.localdate().strftime('%Y%m%d')}-{uuid.uuid4().hex[:5].upper()}"
        
        # البحث عن طريقة دفع بالاسم إن أمكن أو إنشاء معرف افتراضي
        payment_method_id = None
        try:
            from apps.finance.domain.models import PaymentMethod
            pm = PaymentMethod.objects.filter(tenant_id=tenant_id, name_ar__icontains='بنكك').first()
            if not pm:
                pm = PaymentMethod.objects.filter(tenant_id=tenant_id).first()
            if pm:
                payment_method_id = pm.id
        except Exception:
            pass
        if not payment_method_id:
            payment_method_id = uuid.uuid4()

        receipt = Receipt.objects.create(
            tenant_id=tenant_id,
            student_billing_account=installment.student_billing_account,
            receipt_number=receipt_no,
            payment_date=timezone.localdate(),
            amount=amount_to_pay,
            payment_method_id=payment_method_id,
            status='posted'
        )

        return StandardResponse(data={
            'message': 'تم تسجيل السداد بنجاح وإصدار سند القبض.',
            'receipt_number': receipt.receipt_number,
            'installment_id': str(installment.id),
            'installment_status': installment.status,
            'paid_amount': float(installment.paid_amount),
            'remaining_amount': max(0.0, float(installment.amount) - float(installment.paid_amount))
        })

    @action(detail=False, methods=['get'], url_path='export-monthly-dues')
    def export_monthly_dues(self, request):
        """
        تصدير كشف إكسل رسمي باللغة العربية (RTL) للطلاب المستحقين في الأقساط والدفعات لشهر محدد.
        يدعم التصفية حسب السنة، الشهر، وحالة السداد (الكل، المتأخرات فقط، المستحقة، المسددة).
        """
        from django.http import HttpResponse
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        tenant_id = request.tenant_id
        today = timezone.localdate()
        if tenant_id:
            _sync_invoices_to_installments(tenant_id)

        try:
            year = int(request.query_params.get('year', today.year))
            month = int(request.query_params.get('month', today.month))
        except (ValueError, TypeError):
            year, month = today.year, today.month

        status_param = request.query_params.get('status', 'all')

        # الاستعلام الأساسي
        qs = Installment.objects.filter(
            tenant_id=tenant_id,
            due_date__year=year,
            due_date__month=month
        ).select_related('student_billing_account', 'installment_plan', 'invoice')

        if status_param == 'overdue':
            qs = qs.filter(due_date__lt=today).exclude(status='paid')
        elif status_param == 'pending':
            qs = qs.exclude(status='paid')
        elif status_param == 'paid':
            qs = qs.filter(status='paid')

        qs = qs.order_by('due_date', 'student_billing_account__student_id')

        # جلب تفاصيل الطلاب وسريلاتهم
        from apps.student_finance.interfaces.serializers import _extract_student_finance_metadata

        # اسم الشهر بالعربية
        months_ar = {
            1: 'يناير', 2: 'فبراير', 3: 'مارس', 4: 'أبريل',
            5: 'مايو', 6: 'يونيو', 7: 'يوليو', 8: 'أغسطس',
            9: 'سبتمبر', 10: 'أكتوبر', 11: 'نوفمبر', 12: 'ديسمبر'
        }
        month_name = months_ar.get(month, f"شهر {month}")

        # معلومات المؤسسة
        tenant = getattr(request, 'tenant', None)
        if not tenant:
            from apps.tenants.models import Tenant
            tenant = Tenant.objects.filter(id=tenant_id).first() or Tenant.objects.first()
        school_name = (tenant.name_ar or tenant.name) if tenant else 'مدارس المورد النموذجية الخاصة'

        # إنشاء مصنف العمل
        wb = openpyxl.Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet()
        ws.title = f"مستحقي {month_name} {year}"
        ws.views.sheetView[0].rightToLeft = True

        # أنماط التنسيق
        font_title = Font(name='Calibri', size=16, bold=True, color='0F172A')
        font_subtitle = Font(name='Calibri', size=11, bold=True, color='475569')
        font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        font_data = Font(name='Calibri', size=10, color='1E293B')
        font_totals = Font(name='Calibri', size=11, bold=True, color='0F172A')

        fill_header = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        fill_totals = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
        fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')

        thin_side = Side(border_style='thin', color='CBD5E1')
        border_box = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)

        align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        align_right = Alignment(horizontal='right', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')

        # 1. الترويسة الرئيسية
        ws.merge_cells('A1:K1')
        cell_t1 = ws['A1']
        cell_t1.value = school_name
        cell_t1.font = font_title
        cell_t1.alignment = align_center

        ws.merge_cells('A2:K2')
        cell_t2 = ws['A2']
        status_label = 'كافة الأقساط' if status_param == 'all' else ('المتأخرات فقط' if status_param == 'overdue' else ('المستحقة غير المسددة' if status_param == 'pending' else 'المسددة'))
        cell_t2.value = f"كشف استحقاق الأقساط والدفعات الشهرية — {month_name} {year} ({status_label})"
        cell_t2.font = font_subtitle
        cell_t2.alignment = align_center

        ws.merge_cells('A3:K3')
        cell_t3 = ws['A3']
        cell_t3.value = f"تاريخ الاستخراج: {timezone.now().strftime('%Y-%m-%d %H:%M')} | العملة: الجنيه السوداني (ج.س) | إجمالي عدد السجلات: {qs.count()}"
        cell_t3.font = Font(name='Calibri', size=9, italic=True, color='64748B')
        cell_t3.alignment = align_center

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 18

        # 2. رؤوس الأعمدة (الصف 5)
        headers = [
            'م', 'الرقم المدرسي', 'اسم الطالب', 'الصف / المرحلة',
            'ولي الأمر', 'هاتف ولي الأمر', 'تاريخ الاستحقاق',
            'قيمة القسط (ج.س)', 'المسدد (ج.س)', 'المتبقي المستحق (ج.س)', 'حالة القسط'
        ]
        ws.row_dimensions[5].height = 26
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_box

        # 3. ملء البيانات
        row_idx = 6
        tot_amount = 0.0
        tot_paid = 0.0
        tot_rem = 0.0

        for idx, ins in enumerate(qs, start=1):
            ac = ins.student_billing_account
            meta = _extract_student_finance_metadata(ac) if ac else {}
            st_num = meta.get('student_number') or (ac.account_number if ac else '-')
            st_name = meta.get('student_name') or 'طالب'
            grade = (meta.get('grade_name') or '') + (' - ' + meta.get('section_name') if meta.get('section_name') else '')
            g_name = meta.get('guardian_name') or '-'
            g_phone = meta.get('guardian_phone') or '-'

            amt = float(ins.amount or 0.0)
            paid = float(ins.paid_amount or 0.0)
            rem = max(0.0, amt - paid)

            tot_amount += amt
            tot_paid += paid
            tot_rem += rem

            due_str = str(ins.due_date)
            if ins.status == 'paid':
                st_label = 'مسدد بالكامل'
            elif due_str < str(today):
                st_label = 'متأخر السداد'
            else:
                st_label = 'مستحق / مجدول'

            values = [
                idx, st_num, st_name, grade or 'العام الحالي',
                g_name, g_phone, due_str,
                amt, paid, rem, st_label
            ]

            ws.row_dimensions[row_idx].height = 20
            for c_idx, val in enumerate(values, start=1):
                c = ws.cell(row=row_idx, column=c_idx, value=val)
                c.font = font_data
                c.border = border_box

                if c_idx in (1, 2, 6, 7, 11):
                    c.alignment = align_center
                elif c_idx in (8, 9, 10):
                    c.alignment = align_left
                    c.number_format = '#,##0'
                else:
                    c.alignment = align_right

                if idx % 2 == 0:
                    c.fill = fill_zebra

                # تلوين حالة القسط
                if c_idx == 11:
                    if st_label == 'متأخر السداد':
                        c.font = Font(name='Calibri', size=10, bold=True, color='DC2626')
                    elif st_label == 'مسدد بالكامل':
                        c.font = Font(name='Calibri', size=10, bold=True, color='16A34A')
                    else:
                        c.font = Font(name='Calibri', size=10, bold=True, color='D97706')

            row_idx += 1

        # 4. صف الإجماليات
        ws.row_dimensions[row_idx].height = 24
        tot_label_cell = ws.cell(row=row_idx, column=1, value='الإجمالي الكلي المستحق')
        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
        tot_label_cell.font = font_totals
        tot_label_cell.alignment = align_center
        tot_label_cell.fill = fill_totals
        tot_label_cell.border = border_box

        for c_idx in range(1, 8):
            ws.cell(row=row_idx, column=c_idx).border = border_box
            ws.cell(row=row_idx, column=c_idx).fill = fill_totals

        # إجمالي المبالغ
        c_tot_amt = ws.cell(row=row_idx, column=8, value=tot_amount)
        c_tot_amt.font = font_totals
        c_tot_amt.fill = fill_totals
        c_tot_amt.border = border_box
        c_tot_amt.alignment = align_left
        c_tot_amt.number_format = '#,##0'

        c_tot_pd = ws.cell(row=row_idx, column=9, value=tot_paid)
        c_tot_pd.font = Font(name='Calibri', size=11, bold=True, color='16A34A')
        c_tot_pd.fill = fill_totals
        c_tot_pd.border = border_box
        c_tot_pd.alignment = align_left
        c_tot_pd.number_format = '#,##0'

        c_tot_rem = ws.cell(row=row_idx, column=10, value=tot_rem)
        c_tot_rem.font = Font(name='Calibri', size=11, bold=True, color='DC2626')
        c_tot_rem.fill = fill_totals
        c_tot_rem.border = border_box
        c_tot_rem.alignment = align_left
        c_tot_rem.number_format = '#,##0'

        c_end = ws.cell(row=row_idx, column=11, value='ج.س')
        c_end.font = font_totals
        c_end.fill = fill_totals
        c_end.border = border_box
        c_end.alignment = align_center

        # ضبط عروض الأعمدة تلقائياً
        col_widths = {
            1: 6,   # م
            2: 15,  # الرقم المدرسي
            3: 28,  # اسم الطالب
            4: 22,  # الصف
            5: 25,  # ولي الأمر
            6: 16,  # هاتف ولي الأمر
            7: 15,  # تاريخ الاستحقاق
            8: 18,  # قيمة القسط
            9: 18,  # المسدد
            10: 18, # المتبقي
            11: 15  # الحالة
        }
        for col_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # تجهيز الاستجابة كملف Excel قابل للتنزيل
        from io import BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"installments_dues_{year}_{month:02d}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    @action(detail=True, methods=['get'], url_path='reminder-info')
    def reminder_info(self, request, pk=None):
        """تجهيز نص إشعار تذكير بالسداد مناسب للواتساب متوافق مع السياق السوداني."""
        from apps.student_finance.interfaces.serializers import _extract_student_finance_metadata
        installment = self.get_object()
        meta = _extract_student_finance_metadata(installment.student_billing_account)
        student_name = meta.get('student_name') or 'الطالب'
        guardian_phone = meta.get('guardian_phone') or ''
        remaining = max(0.0, float(installment.amount) - float(installment.paid_amount))
        formatted_amount = f"{remaining:,.0f} ج.س"
        
        clean_phone = guardian_phone.replace(' ', '').replace('-', '').replace('+', '')
        if clean_phone.startswith('0'):
            clean_phone = '249' + clean_phone[1:]
        elif not clean_phone.startswith('249') and len(clean_phone) == 9:
            clean_phone = '249' + clean_phone

        text = (
            f"السلام عليكم ورحمة الله وبركاته،\n"
            f"الأخ/الأخت ولي أمر الطالب: {student_name}\n"
            f"نود تذكيركم بموعد استحقاق القسط المدرسي بمبلغ ({formatted_amount}) والمستحق بتاريخ {installment.due_date}.\n\n"
            f"طرق السداد المعتمدة لدى المدرسة:\n"
            f"• تطبيق بنكك (بنك الخرطوم)\n"
            f"• تطبيق فوري (بنك فيصل الإسلامي)\n"
            f"• أوكاش (بنك أمدرمان الوطني)\n\n"
            f"الرجاء إرسال إشعار التحويل البنكي مع ذكر اسم الطالب لتحديث السجل المالي فوراً.\n"
            f"شاكرين لكم حسن تعاونكم الدائم.\n"
            f"إدارة الحسابات المدرسية — منظومة نبراس التعليمية"
        )

        import urllib.parse
        encoded_text = urllib.parse.quote(text)
        whatsapp_url = f"https://wa.me/{clean_phone}?text={encoded_text}" if clean_phone else ""

        return StandardResponse(data={
            'student_name': student_name,
            'guardian_phone': guardian_phone,
            'clean_phone': clean_phone,
            'amount': remaining,
            'formatted_amount': formatted_amount,
            'due_date': str(installment.due_date),
            'reminder_text': text,
            'whatsapp_url': whatsapp_url
        })


class StudentReceivableViewSet(BaseCRUDViewSet):
    model_class = StudentReceivable
    serializer_class = StudentReceivableSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        account = self.request.query_params.get('student_billing_account')
        status_param = self.request.query_params.get('status')
        if account:
            qs = qs.filter(student_billing_account_id=account)
        if status_param:
            qs = qs.filter(status=status_param)
        return qs


class PaymentAllocationViewSet(BaseCRUDViewSet):
    model_class = PaymentAllocation
    serializer_class = PaymentAllocationSerializer


class ReceiptViewSet(BaseCRUDViewSet):
    model_class = Receipt
    serializer_class = ReceiptSerializer
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ['receipt_number']

    def get_queryset(self):
        qs = super().get_queryset()
        account = self.request.query_params.get('student_billing_account')
        if account:
            qs = qs.filter(student_billing_account_id=account)
        return qs

    @action(detail=False, methods=['post'], url_path='receive-payment')
    def receive_payment(self, request):
        tenant_id = request.tenant_id
        billing_account_id = request.data.get('billing_account_id')
        amount = request.data.get('amount')
        payment_method_id = request.data.get('payment_method_id')
        bank_account_id = request.data.get('bank_account_id')
        cash_box_id = request.data.get('cash_box_id')

        if not billing_account_id or amount is None or not payment_method_id:
            return Response({'error': 'billing_account_id, amount, and payment_method_id are required'}, status=status.HTTP_400_BAD_REQUEST)

        receipt = PaymentService.receive_payment(
            tenant_id=tenant_id,
            billing_account_id=billing_account_id,
            amount=amount,
            payment_method_id=payment_method_id,
            bank_account_id=bank_account_id,
            cash_box_id=cash_box_id,
            user_id=request.user.id if request.user else None
        )
        serializer = self.get_serializer(receipt)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class RefundViewSet(BaseCRUDViewSet):
    model_class = Refund
    serializer_class = RefundSerializer


class CreditNoteViewSet(BaseCRUDViewSet):
    model_class = CreditNote
    serializer_class = CreditNoteSerializer


class DebitNoteViewSet(BaseCRUDViewSet):
    model_class = DebitNote
    serializer_class = DebitNoteSerializer


class LateFeeRuleViewSet(BaseCRUDViewSet):
    model_class = LateFeeRule
    serializer_class = LateFeeRuleSerializer


class CollectionPolicyViewSet(BaseCRUDViewSet):
    model_class = CollectionPolicy
    serializer_class = CollectionPolicySerializer


class FinancialHoldViewSet(BaseCRUDViewSet):
    model_class = FinancialHold
    serializer_class = FinancialHoldSerializer

    def get_queryset(self):
        qs = super().get_queryset()
        account = self.request.query_params.get('student_billing_account')
        status_param = self.request.query_params.get('status')
        if account:
            qs = qs.filter(student_billing_account_id=account)
        if status_param:
            qs = qs.filter(status=status_param)
        return qs

    @action(detail=True, methods=['post'], url_path='release')
    def release_hold(self, request, pk=None):
        hold = self.get_queryset().get(id=pk)
        hold.status = 'released'
        hold.released_at = timezone.now()
        hold.save(update_fields=['status', 'released_at'])
        # رفع علم الحظر عن الحساب إذا لم تبقَ حالات حظر نشطة
        acc = hold.student_billing_account
        if not FinancialHold.objects.filter(student_billing_account=acc, status='active').exists():
            acc.financial_hold = False
            acc.save(update_fields=['financial_hold'])
        return StandardResponse(data=self.get_serializer(hold).data, message="تم رفع الحظر المالي بنجاح.")

    @action(detail=False, methods=['post'], url_path='apply-hold')
    def apply_hold(self, request):
        tenant_id = request.tenant_id
        billing_account_id = request.data.get('billing_account_id')
        hold_type = request.data.get('hold_type')
        reason = request.data.get('reason')

        hold = HoldService.apply_financial_hold(
            tenant_id=tenant_id,
            billing_account_id=billing_account_id,
            hold_type=hold_type,
            reason=reason,
            user_id=request.user.id if request.user else None
        )
        serializer = self.get_serializer(hold)
        return Response(serializer.data, status=status.HTTP_201_CREATED)


class BillingCycleViewSet(BaseCRUDViewSet):
    model_class = BillingCycle
    serializer_class = BillingCycleSerializer


class StatementViewSet(BaseCRUDViewSet):
    model_class = Statement
    serializer_class = StatementSerializer


class BillingAuditViewSet(BaseCRUDViewSet):
    model_class = BillingAudit
    serializer_class = BillingAuditSerializer


class StudentFinanceSettingsViewSet(BaseCRUDViewSet):
    model_class = StudentFinanceSettings
    serializer_class = StudentFinanceSettingsSerializer


class OnlinePaymentRequestViewSet(BaseCRUDViewSet):
    """طلبات السداد الأونلاين لأولياء الأمور (تحويل بنكي + مراجعة المحاسب)."""
    model_class = OnlinePaymentRequest
    serializer_class = OnlinePaymentRequestSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    filter_backends = [filters.OrderingFilter]
    ordering = ['-created_at']

    def _tenant(self, request):
        return request.tenant.id if hasattr(request, 'tenant') and request.tenant else None

    def get_queryset(self):
        qs = super().get_queryset()
        status_param = self.request.query_params.get('status')
        account = self.request.query_params.get('student_billing_account')
        student = self.request.query_params.get('student_id')
        mine = self.request.query_params.get('mine')
        if status_param:
            qs = qs.filter(status=status_param)
        if account:
            qs = qs.filter(student_billing_account_id=account)
        if student:
            qs = qs.filter(student_id=student)
        if mine in ('1', 'true', 'yes') and self.request.user:
            qs = qs.filter(submitted_by_user_id=self.request.user.id)
        return qs

    def create(self, request, *args, **kwargs):
        """تقديم طلب سداد جديد من ولي الأمر (multipart مع إيصال التحويل)."""
        tenant_id = self._tenant(request)
        data = request.data
        try:
            req = online_payment_service.submit_payment_request(
                tenant_id=tenant_id,
                billing_account_id=data.get('student_billing_account') or data.get('billing_account_id'),
                amount=data.get('amount'),
                transfer_reference=data.get('transfer_reference'),
                transfer_date=data.get('transfer_date'),
                receipt_attachment=request.FILES.get('receipt_attachment'),
                sender_name=data.get('sender_name'),
                note=data.get('note'),
                bank_name=data.get('bank_name', 'بنك الخرطوم'),
                submitted_by_user_id=request.user.id if request.user else None,
            )
        except DjangoValidationError as e:
            return StandardResponse(
                message=str(e.message if hasattr(e, 'message') else e),
                success=False,
                status=status.HTTP_400_BAD_REQUEST
            )
        serializer = self.get_serializer(req)
        return StandardResponse(serializer.data,
                                message="تم استلام طلب السداد وهو الآن قيد مراجعة المحاسبة.",
                                status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['post'], url_path='approve')
    def approve(self, request, pk=None):
        tenant_id = self._tenant(request)
        try:
            req = online_payment_service.approve_payment_request(
                tenant_id=tenant_id, request_id=pk,
                reviewer_id=request.user.id if request.user else None,
            )
        except DjangoValidationError as e:
            return Response({'error': str(e.message if hasattr(e, 'message') else e)},
                            status=status.HTTP_400_BAD_REQUEST)
        return StandardResponse(self.get_serializer(req).data,
                                message="تم اعتماد السداد وتحديث حساب الطالب المالي.")

    @action(detail=True, methods=['post'], url_path='reject')
    def reject(self, request, pk=None):
        tenant_id = self._tenant(request)
        try:
            req = online_payment_service.reject_payment_request(
                tenant_id=tenant_id, request_id=pk,
                reviewer_id=request.user.id if request.user else None,
                reason=request.data.get('reason'),
            )
        except DjangoValidationError as e:
            return Response({'error': str(e.message if hasattr(e, 'message') else e)},
                            status=status.HTTP_400_BAD_REQUEST)
        return StandardResponse(self.get_serializer(req).data,
                                message="تم رفض طلب السداد وإشعار ولي الأمر.")

    @action(detail=False, methods=['get'], url_path='pending-count')
    def pending_count(self, request):
        n = self.get_queryset().filter(status='pending').count()
        return StandardResponse({'pending': n}, message="عدد الطلبات المعلقة.")
